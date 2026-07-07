from openpyxl import Workbook
from openpyxl.reader.excel import load_workbook

workbook=Workbook()
workbook.save("text.xlsx")
print("Excel created successfully")

workbook = Workbook()      # Create an instance
sheet = workbook.active    # Get the active worksheet
sheet.title = "employees"
workbook.save("employees.xlsx")

sheet["A1"]="Name"
sheet["B1"]="Department"

sheet["A2"]="Vinu"
sheet["B2"]="QA"

workbook.save("employees.xlsx")
print("Excel  successfully")

#read excel
workbook=load_workbook("employees.xlsx")
sheet = workbook.active
print(sheet["A2"].value)

#read multiple cells
print(sheet["A2"].value)
print(sheet["B2"].value)

#update cell
sheet["B2"]="Automation"
workbook.save("employees.xlsx")

#read using Row & column
print(sheet.cell(row=2,column=2).value)

#write using row and coloumn
sheet.cell(row=3,column=1).value="Rahul"
sheet.cell(row=3,column=2).value="developer"
workbook.save("employees.xlsx")
print("done")

#find max column
print(sheet.max_column)

#read entire excel
for row in sheet.iter_rows(values_only=True):
    print(row)

#loop through rows
for row in range(2,sheet.max_row+1):
    name=sheet.cell(row=row,column=1).value
    dept=sheet.cell(row=row,column=2).value
    print(name,dept)

#add new sheet
workbook.create_sheet("reports")
workbook.save("employees.xlsx")

#rename sheet
sheet.title="EmployeeData"

workbook.remove(workbook["reports"])
